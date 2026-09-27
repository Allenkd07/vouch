"""Render a ResumeDoc to PDF (Typst) and DOCX (python-docx)."""

import json
import shutil
import tempfile
from pathlib import Path

import pymupdf
import typst
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from vouch.tailoring.document import ResumeDoc

TEMPLATE = Path(__file__).resolve().parents[3] / "resume_templates" / "resume.typ"


def contact_lines(doc: ResumeDoc) -> list[str]:
    """Contact details, then links: two short lines instead of one that wraps mid-URL."""
    c = doc.contact
    links = [link.url.removeprefix("https://").removeprefix("http://") for link in c.links]
    lines = [" | ".join(p for p in (c.email, c.phone, c.location) if p), " | ".join(links)]
    return [line for line in lines if line]


def render_pdf(doc: ResumeDoc, template: Path = TEMPLATE) -> bytes:
    data = doc.model_dump(mode="json")
    data["contact_lines"] = contact_lines(doc)
    with tempfile.TemporaryDirectory() as tmp:
        shutil.copy(template, Path(tmp) / "resume.typ")
        (Path(tmp) / "resume.json").write_text(json.dumps(data), encoding="utf-8")
        return typst.compile(str(Path(tmp) / "resume.typ"))


def page_count(pdf: bytes) -> int:
    with pymupdf.open(stream=pdf, filetype="pdf") as d:
        return d.page_count


def render_docx(doc: ResumeDoc, path: Path) -> None:
    d = Document()
    style = d.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10)
    for section in d.sections:
        section.left_margin = section.right_margin = Pt(42)
        section.top_margin = section.bottom_margin = Pt(34)

    def para(text: str = "", bold=False, size=None, align=None, italic=False):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if align:
            p.alignment = align
        return p

    def heading(title: str):
        para(title.upper(), bold=True, size=11).paragraph_format.space_before = Pt(6)

    def entry(e):
        p = para(e.heading, bold=True)
        if e.subheading:
            p.add_run(f" | {e.subheading}")
        if e.dates:
            p.add_run(f"    {e.dates}")
        extra = " | ".join(x for x in (e.description, e.location) if x)
        if extra:
            para(extra, italic=True, size=9.5)
        for b in e.bullets:
            bp = d.add_paragraph(b.text, style="List Bullet")
            bp.paragraph_format.space_after = Pt(1)

    para(doc.contact.name, bold=True, size=17, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in contact_lines(doc):
        para(line, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    if doc.summary:
        heading("Summary")
        para(doc.summary)
    heading("Skills")
    for s in doc.skills:
        p = para(f"{s.label}: ", bold=True)
        p.add_run(", ".join(s.items))
    heading("Experience")
    for e in doc.experience:
        entry(e)
    if doc.projects:
        heading("Projects")
        for e in doc.projects:
            entry(e)
    heading("Education")
    for e in doc.education:
        entry(e)
    d.save(str(path))
